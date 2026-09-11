# -*- coding: utf-8 -*-
"""docx → 뉴스레터용 마크다운 + 이미지 추출.

build.py(노션 마크다운 기준)가 읽을 수 있는 형태로 변환한다.
- 문단/표를 문서 순서대로 순회
- 인라인 이미지는 추출해 ![](파일명) 으로 치환
- 굵게는 **...** 로 보존
"""
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def rel_image_map(doc):
    """rId → (부품이름, 바이트)"""
    out = {}
    for rid, rel in doc.part.rels.items():
        if "image" in rel.reltype:
            try:
                out[rid] = (Path(rel.target_part.partname).name,
                            rel.target_part.blob)
            except Exception:
                pass
    return out


def para_images(p):
    """문단 안의 이미지 rId 목록 (문서 순서)"""
    rids = []
    for blip in p._p.iter(f"{A}blip"):
        rid = blip.get(f"{R}embed") or blip.get(f"{R}link")
        if rid:
            rids.append(rid)
    # VML(구형) 이미지도 훑는다
    for imagedata in p._p.iter(
            "{urn:schemas-microsoft-com:vml}imagedata"):
        rid = imagedata.get(f"{R}id")
        if rid:
            rids.append(rid)
    return rids


def run_text(r):
    t = "".join(n.text or "" for n in r._element.iter(f"{W}t"))
    if not t:
        return ""
    if r.bold:
        return f"**{t}**"
    return t


def para_text(p):
    s = "".join(run_text(r) for r in p.runs)
    # 연속 볼드 마커 정리: **A****B** → **AB**
    s = s.replace("****", "")
    return s.strip()


def convert(docx_path, out_dir, slug, img_prefix=None):
    docx_path = Path(docx_path)
    out_dir = Path(out_dir)
    img_dir = out_dir / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(str(docx_path))
    imgs = rel_image_map(doc)
    img_prefix = img_prefix or slug

    saved = {}          # rId → 상대경로
    counter = [0]

    def save_img(rid, hint=""):
        if rid in saved:
            return saved[rid]
        if rid not in imgs:
            return None
        name, blob = imgs[rid]
        ext = Path(name).suffix.lower() or ".png"
        counter[0] += 1
        suffix = hint or f"{counter[0]:02d}"
        fname = f"{img_prefix}_{suffix}{ext}"
        (img_dir / fname).write_bytes(blob)
        rel = f"images/{fname}"
        saved[rid] = rel
        return rel

    lines = []
    body = doc.element.body
    # 문단·표를 문서 순서대로
    para_iter = iter(doc.paragraphs)
    tbl_iter = iter(doc.tables)

    for child in body.iterchildren():
        if child.tag == f"{W}p":
            p = next(para_iter, None)
            if p is None:
                continue
            rids = para_images(p)
            txt = para_text(p)
            style = (p.style.name or "").lower()

            if rids:
                for rid in rids:
                    rel = save_img(rid)
                    if rel:
                        lines.append(f"![image]({rel})")
                        lines.append("")
                if txt:
                    lines.append(txt)
                    lines.append("")
                continue

            if not txt:
                lines.append("")
                continue

            if style.startswith("heading 1") or style.startswith("제목 1"):
                lines.append(f"# {txt}")
            elif style.startswith("heading 2") or style.startswith("제목 2"):
                lines.append(f"## {txt}")
            elif style.startswith("heading 3"):
                lines.append(f"### {txt}")
            elif style.startswith("list") or style.startswith("목록"):
                lines.append(f"- {txt}")
            else:
                lines.append(txt)
            lines.append("")

        elif child.tag == f"{W}tbl":
            t = next(tbl_iter, None)
            if t is None:
                continue
            rows = []
            for row in t.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                rows.append(cells)
            if not rows:
                continue
            # 표 안에만 이미지가 있는 경우도 건진다
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for rid in para_images(p):
                            rel = save_img(rid)
                            if rel:
                                lines.append(f"![image]({rel})")
                                lines.append("")
            ncol = max(len(r) for r in rows)
            head = rows[0] + [""] * (ncol - len(rows[0]))
            lines.append("| " + " | ".join(head) + " |")
            lines.append("|" + "---|" * ncol)
            for r in rows[1:]:
                r = r + [""] * (ncol - len(r))
                lines.append("| " + " | ".join(r) + " |")
            lines.append("")

    md = "\n".join(lines)
    md = re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"
    return md, sorted(saved.values())


if __name__ == "__main__":
    src, out, slug = sys.argv[1], sys.argv[2], sys.argv[3]
    md, images = convert(src, out, slug)
    Path(out).mkdir(parents=True, exist_ok=True)
    (Path(out) / f"{slug}.md").write_text(md, encoding="utf-8")
    print(f"{slug}: {len(md)}자, 이미지 {len(images)}개")
